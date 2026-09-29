# -*- coding: utf-8 -*-
"""Completa el catálogo con lo que la app ya sabía y el catálogo no.

Qué hace, y por qué cada cosa:

  1. RELLENA seis `Cadena principal` vacías. Son huecos, no criterio: un Spoto
     Press es empuje horizontal y un KB Bottoms-Up Press es vertical, no hay
     nada que discutir. Salieron de cruzar los pools de la app contra el
     catálogo — la app los tenía clasificados, el catálogo los tenía en blanco.

  2. Agrega `traccion vertical` como CADENA SECUNDARIA a los tres High Pull,
     dejando `bisagra` como principal. Decisión suya, con estas palabras: "las
     dos cosas". Arrancan del suelo con cadera (bisagra) y el codo tira alto
     (tracción vertical); el catálogo tiene dos columnas justamente para esto.

  3. Corrige `Ring Row vertical (pies elevados)` a `traccion vertical`. También
     decisión suya: en la app está puesto como progresión hacia el pull-up, y
     ese es el sentido que tiene en un pool de tracción vertical.

  4. Escribe la hoja `Pools` con las 228 filas de GEN_POOLS. NO es una columna
     del catálogo y no puede serlo: 31 de los 196 movimientos están en más de
     un pool —Back Squat es fuerza/squat y musculacion/legs— así que la
     relación es de muchos a muchos y necesita su propia hoja.

Lo que NO hace, y es lo importante: no genera los pools ni los va a generar.
El catálogo ofrece entre 7 y 20 candidatos por pool y la app elige seis. Esa
elección de seis es curaduría suya y no está en ninguna columna; no se calcula.
El catálogo VERIFICA los pools (ver revisar_pools.py), no los produce.

Es idempotente: correrlo dos veces deja el archivo igual. Y avisa fila por fila
lo que cambió, porque un catálogo que se edita en silencio es peor que uno
incompleto.

    python3 enriquecer_catalogo.py [--escribir]

Sin --escribir solo dice qué haría.
"""
import sys

import openpyxl

from pools import leer_pools, por_movimiento

CATALOGO = 'CSL-Fit_Catalogo_Completo.xlsx'
HOJA = 'Movimientos'
HOJA_POOLS = 'Pools'

COL_NOMBRE, COL_BASE, COL_PATRON, COL_CAD1, COL_CAD2 = 1, 2, 3, 4, 5

# 5. Cinco zancadas que el catálogo dejó en '(sin clasificar)'. No es criterio:
# las cinco ya tienen `Cadena principal = rodilla` y las cinco están en el pool
# `fuerza/lunge` de la app. Solo faltaba el patrón. El planificador no lee la
# columna Patrón —lee Categoría y las cadenas— así que esto no puede cambiar
# un WOD; lo arregla para el checker y para quien abra la planilla.
PATRON = {
    'Barbell Reverse Lunge':    'lunge',
    'DB Bulgarian Split Squat': 'lunge',
    'DB Lateral Lunge':         'lunge',
    'Front Rack Step-up':       'lunge',
    'Walking Lunge (barra)':    'lunge',
}

# 1. Los seis huecos. Salieron del cruce, no de mi opinión.
CADENA_PRINCIPAL = {
    'Behind-the-Neck Press':            'empuje vertical',
    'DB Z Press (sentado en el suelo)': 'empuje vertical',
    'KB Bottoms-Up Press':              'empuje vertical',
    'DB Incline Press (pies en banco)': 'empuje horizontal',
    'DB Squeeze Press':                 'empuje horizontal',
    'Spoto Press':                      'empuje horizontal',
    # 3. Decisión suya: acá manda la app.
    'Ring Row vertical (pies elevados)': 'traccion vertical',
}

# 2. Decisión suya: "las dos cosas". La principal se queda como está.
CADENA_SECUNDARIA = {
    'Clean High Pull': 'traccion vertical',
    'DB High Pull':    'traccion vertical',
    'KB High Pull':    'traccion vertical',
}


def enriquecer(ruta=CATALOGO, escribir=False):
    wb = openpyxl.load_workbook(ruta)
    ws = wb[HOJA]
    cambios = []

    for fila in ws.iter_rows(min_row=2):
        nombre = str(fila[0].value).strip() if fila[0].value else ''
        if not nombre:
            continue
        for columna, tabla, etiqueta in (
                (COL_PATRON, PATRON, 'Patrón'),
                (COL_CAD1, CADENA_PRINCIPAL, 'Cadena principal'),
                (COL_CAD2, CADENA_SECUNDARIA, 'Cadena secundaria')):
            if nombre not in tabla:
                continue
            celda = ws.cell(fila[0].row, columna)
            nuevo = tabla[nombre]
            if celda.value == nuevo:
                continue
            cambios.append((fila[0].row, nombre, etiqueta, celda.value, nuevo))
            if escribir:
                celda.value = nuevo

    pools = leer_pools()
    donde = por_movimiento(pools)
    filas_pools = sorted(
        (m, foco, clave, equipo)
        for m, ds in donde.items() for foco, clave, equipo in ds)

    if escribir:
        if HOJA_POOLS in wb.sheetnames:
            del wb[HOJA_POOLS]
        hp = wb.create_sheet(HOJA_POOLS)
        hp.append(['Movimiento', 'Foco', 'Clave del pool', 'Equipo'])
        for f in filas_pools:
            hp.append(list(f))
        hp.freeze_panes = 'A2'
        wb.save(ruta)

    return cambios, filas_pools


if __name__ == '__main__':
    escribir = '--escribir' in sys.argv
    cambios, filas = enriquecer(escribir=escribir)
    if not cambios:
        print('las celdas ya estaban como corresponde — nada que cambiar')
    for row, nombre, etiqueta, antes, nuevo in cambios:
        print(f'  fila {row:>4}  {nombre:36} {etiqueta:18} '
              f'{antes!r} -> {nuevo!r}')
    print(f'\nhoja «{HOJA_POOLS}»: {len(filas)} filas '
          f'({len({f[0] for f in filas})} movimientos distintos)')
    print('ESCRITO en ' + CATALOGO if escribir
          else 'nada escrito — volvé a correr con --escribir')
